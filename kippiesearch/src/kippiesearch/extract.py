"""Tekstextractie uit lokale bestanden. Afbeeldingen worden overgeslagen (OCR uit)."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from kippiesearch.logging_setup import get_logger

log = get_logger("extract")

TEXT_EXT = {
    ".txt",
    ".md",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".org",
    ".tex",
    ".rtf",
}
CODE_EXT = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".cc",
    ".cpp",
    ".h",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".sh",
    ".bash",
    ".zsh",
    ".ps1",
    ".yml",
    ".yaml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".sql",
    ".css",
    ".scss",
    ".kt",
    ".swift",
    ".r",
    ".lua",
    ".pl",
    ".pm",
    ".m",
    ".mm",
    ".scala",
    ".dart",
    ".vue",
    ".svelte",
    ".xml",
    ".gradle",
    ".cmake",
    ".make",
    ".mk",
    ".dockerfile",
}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
PPTX_EXT = {".pptx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp", ".heic"}

SUPPORTED_EXT = TEXT_EXT | CODE_EXT | PDF_EXT | DOCX_EXT | PPTX_EXT


def file_type_for(path: Path | str) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in PDF_EXT:
        return "pdf"
    if suffix in DOCX_EXT:
        return "docx"
    if suffix in PPTX_EXT:
        return "pptx"
    if suffix in CODE_EXT:
        return "code"
    if suffix in IMAGE_EXT:
        return "image"
    if suffix in TEXT_EXT:
        return "text"
    return suffix.lstrip(".") or "unknown"


def is_supported(path: Path | str, ocr_enabled: bool = False) -> bool:
    suffix = Path(path).suffix.lower()
    if suffix in SUPPORTED_EXT:
        return True
    return bool(ocr_enabled and suffix in IMAGE_EXT)


def _decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: Path, ocr_enabled: bool = False) -> str:
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXT:
        if not ocr_enabled:
            return ""
        log.info("ocr_skipped_not_installed", path=str(path))
        return ""
    if suffix in PDF_EXT:
        return _extract_pdf(path)
    if suffix in DOCX_EXT:
        return _extract_docx(path)
    if suffix in PPTX_EXT:
        return _extract_pptx(path)
    if suffix in {".html", ".htm"}:
        return _extract_html(path)
    if suffix == ".json":
        return _extract_json(path)
    return _extract_plain(path)


def extract_from_bytes(name: str, data: bytes, ocr_enabled: bool = False) -> str:
    suffix = Path(name).suffix.lower()
    tmp_hint = Path(name)
    if suffix in IMAGE_EXT:
        return "" if not ocr_enabled else ""
    if suffix in PDF_EXT:
        return _extract_pdf_bytes(data)
    if suffix in {".html", ".htm"}:
        return _strip_html(_decode_bytes(data))
    if suffix == ".json":
        return _json_to_text(_decode_bytes(data))
    if suffix in TEXT_EXT | CODE_EXT:
        return _decode_bytes(data)
    if suffix in DOCX_EXT | PPTX_EXT:
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=suffix, delete=True) as handle:
            handle.write(data)
            handle.flush()
            return extract_text(Path(handle.name), ocr_enabled=ocr_enabled)
    _ = tmp_hint
    return _decode_bytes(data)


def _extract_plain(path: Path) -> str:
    return _decode_bytes(path.read_bytes())


def _extract_html(path: Path) -> str:
    return _strip_html(_extract_plain(path))


def _strip_html(html: str) -> str:
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "lxml")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return soup.get_text("\n", strip=True)
    except Exception:
        return html


def _extract_json(path: Path) -> str:
    return _json_to_text(_extract_plain(path))


def _json_to_text(raw: str) -> str:
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    return json.dumps(parsed, ensure_ascii=False, indent=2)


def _extract_pdf(path: Path) -> str:
    return _extract_pdf_bytes(path.read_bytes())


def _extract_pdf_bytes(data: bytes) -> str:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            log.warning("pdf_page_failed", error=str(exc))
            text = ""
        if text.strip():
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


def iter_supported(paths: Iterable[Path], ocr_enabled: bool = False) -> Iterable[Path]:
    for path in paths:
        if is_supported(path, ocr_enabled=ocr_enabled):
            yield path
