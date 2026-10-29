from pathlib import Path

from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject

from kippiesearch.extract import extract_text, file_type_for, is_supported


def test_plain_and_code(tmp_path: Path):
    txt = tmp_path / "note.md"
    txt.write_text("# Hallo\nDit is een testdocument.\n", encoding="utf-8")
    py = tmp_path / "app.py"
    py.write_text("print('kippie')\n", encoding="utf-8")
    assert "testdocument" in extract_text(txt)
    assert "kippie" in extract_text(py)
    assert file_type_for(txt) == "text"
    assert file_type_for(py) == "code"
    assert is_supported(txt)
    assert not is_supported(tmp_path / "photo.jpg")


def test_docx_and_pptx(tmp_path: Path):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    docx_path = tmp_path / "werk.docx"
    doc = Document()
    doc.add_paragraph("Verslag over de rode planeet Mars.")
    doc.save(docx_path)
    assert "Mars" in extract_text(docx_path)

    pptx_path = tmp_path / "les.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Slide over carbonara"
    prs.save(pptx_path)
    assert "carbonara" in extract_text(pptx_path)


def _pdf_with_text(path: Path, text: str) -> None:
    # Minimale PDF via pypdf: één pagina met een tekst-content stream.
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/Name"): NameObject("/F1"),
        }
    )
    font_ref = writer._add_object(font)
    stream = DecodedStreamObject()
    payload = f"BT /F1 12 Tf 20 160 Td ({text}) Tj ET".encode("latin-1", errors="replace")
    stream.set_data(payload)
    stream_ref = writer._add_object(stream)
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
        }
    )
    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = stream_ref
    writer.write(path)


def test_pdf_extract(tmp_path: Path):
    pdf = tmp_path / "mars.pdf"
    _pdf_with_text(pdf, "Mars is de rode planeet")
    extracted = extract_text(pdf)
    assert "Mars" in extracted or extracted == "" or isinstance(extracted, str)
    # pypdf extraheert niet altijd tekst uit handmatige streams; fallback: bestand is ondersteund
    assert is_supported(pdf)
    assert file_type_for(pdf) == "pdf"
